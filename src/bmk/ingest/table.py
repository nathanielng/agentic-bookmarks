"""CSV and XLSX (PLAN03 §5.5).

The URL column is a header matching url/link/href/website/address, otherwise the first column whose values
are more than 80% URL-like. `url_col`, `title_col` and `date_col` override detection (a header name or a
1-based column number). When the URL column can't be chosen, `ColumnsNeeded` carries the columns and a
10-row preview so the caller (CLI or UI picker) can ask. XLSX reads the cell's hyperlink target as well as
its value, since the displayed text is often not the link.

Locators are `row 42` (CSV) or `Sheet1!R42` (XLSX). Unmapped columns become remark lines on new bookmarks.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any
from zoneinfo import ZoneInfo

from ..models import RawItem
from .links import clean_title, looks_like_url, with_scheme
from .rtf import FileFormatError

_URL_HEADER = re.compile(r"(?i)\b(url|link|href|website|address)s?\b")
_TITLE_HEADER = re.compile(r"(?i)\b(title|name|description|text|label)\b")
_DATE_HEADER = re.compile(r"(?i)\b(date|time|saved|created|added|timestamp)\b")
PREVIEW_ROWS = 10


@dataclass
class Columns:
    url: int
    title: int | None = None
    date: int | None = None
    header: bool = True  # first row is a header


@dataclass
class Grid:
    """One sheet: rows of display values plus, per cell, a hyperlink target if it has one."""

    name: str | None  # sheet name (XLSX) or None (CSV)
    rows: list[list[Any]]
    links: dict[tuple[int, int], str] = field(default_factory=dict)


class ColumnsNeeded(FileFormatError):
    def __init__(self, sheet: str | None, headers: list[str], preview: list[list[str]], reason: str):
        self.sheet, self.headers, self.preview = sheet, headers, preview
        super().__init__(f"{reason}; choose the URL column with --url-col (columns: {', '.join(headers)})")

    def as_dict(self) -> dict[str, Any]:
        return {"sheet": self.sheet, "headers": self.headers, "preview": self.preview, "error": str(self)}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value).strip()


def _url_like(grid: Grid, r: int, c: int) -> bool:
    return (r, c) in grid.links or looks_like_url(_text(grid.rows[r][c]) if c < len(grid.rows[r]) else "")


def _headers(grid: Grid, header: bool) -> list[str]:
    width = max((len(r) for r in grid.rows), default=0)
    first = grid.rows[0] if grid.rows and header else []
    return [_text(first[c]) if c < len(first) and _text(first[c]) else f"column {c + 1}" for c in range(width)]


def _pick(spec: str | None, headers: list[str]) -> int | None:
    if spec is None or spec == "":
        return None
    if str(spec).isdigit() and 1 <= int(spec) <= len(headers):
        return int(spec) - 1
    lowered = [h.lower() for h in headers]
    if str(spec).lower() in lowered:
        return lowered.index(str(spec).lower())
    raise FileFormatError(f"no column '{spec}' (columns: {', '.join(headers)})")


def detect(grid: Grid, url_col=None, title_col=None, date_col=None) -> Columns:
    if not grid.rows:
        raise FileFormatError("no rows")
    width = max(len(r) for r in grid.rows)
    first = [_text(v) for v in grid.rows[0]]
    # A header row has no URL-like cell; a file of bare URLs has none.
    header = not any(_url_like(grid, 0, c) for c in range(len(first))) and any(first)
    headers = _headers(grid, header)
    body = range(1 if header else 0, len(grid.rows))

    url = _pick(url_col, headers)
    if url is None:
        named = [c for c, h in enumerate(headers) if header and _URL_HEADER.search(h)]
        if len(named) > 1:  # e.g. "URL" and "Website": prefer the one that holds URLs
            named = [c for c in named if _ratio(grid, body, c) > 0.8] or named
        if len(named) == 1:
            url = named[0]
        elif not named:
            url = next((c for c in range(width) if _ratio(grid, body, c) > 0.8), None)
        if url is None:
            reason = "More than one URL column" if named else "No URL column found"
            preview = [[_text(v) for v in row] for row in grid.rows[(1 if header else 0) :][:PREVIEW_ROWS]]
            raise ColumnsNeeded(grid.name, headers, preview, reason)

    def by_header(pattern):
        return next((c for c, h in enumerate(headers) if header and c != url and pattern.search(h)), None)

    title = _pick(title_col, headers) if title_col else by_header(_TITLE_HEADER)
    when = _pick(date_col, headers) if date_col else by_header(_DATE_HEADER)
    return Columns(url=url, title=title, date=when, header=header)


def _ratio(grid: Grid, body: range, c: int) -> float:
    values = [r for r in body if c < len(grid.rows[r]) and (_text(grid.rows[r][c]) or (r, c) in grid.links)]
    return sum(_url_like(grid, r, c) for r in values) / len(values) if values else 0.0


def _date(value: Any, tz: str) -> datetime | None:
    zone = ZoneInfo(tz)
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=zone)
    if isinstance(value, date):
        return datetime.combine(value, time(), zone)
    text = _text(value)
    if not text:
        return None
    try:
        dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y/%m/%d %H:%M", "%Y/%m/%d", "%d %b %Y", "%b %d, %Y"):
            try:
                dt = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        else:
            return None
    return dt if dt.tzinfo else dt.replace(tzinfo=zone)


def rows_to_items(grid: Grid, cols: Columns, tz: str) -> tuple[list[RawItem], list[str]]:
    headers = _headers(grid, cols.header)
    items, warnings = [], []
    for r in range(1 if cols.header else 0, len(grid.rows)):
        row = grid.rows[r]
        if not any(_text(v) for v in row) and not any((r, c) in grid.links for c in range(len(row))):
            continue  # blank row
        cell = lambda c: row[c] if c is not None and c < len(row) else None  # noqa: E731
        shown = _text(cell(cols.url))
        target = grid.links.get((r, cols.url))
        url = with_scheme(target or shown)
        title = _text(cell(cols.title)) if cols.title is not None else None
        if not title and target and shown and not looks_like_url(shown):
            title = shown  # a hyperlinked cell: its text is the title
        item = RawItem(url=url, title=clean_title(title), locator=f"{grid.name}!R{r + 1}" if grid.name else f"row {r + 1}")
        if cols.date is not None and _text(cell(cols.date)):
            if stamp := _date(cell(cols.date), tz):
                item.event_at, item.event_at_kind = stamp, "exact"
            else:
                warnings.append(f"{item.locator}: unreadable date '{_text(cell(cols.date))}'")
        for c, h in enumerate(headers):
            if c not in (cols.url, cols.title, cols.date) and (value := _text(cell(c))):
                item.extra[h] = value
        items.append(item)
    return items, warnings


# ---------------------------------------------------------------------------------------------- readers


def read_csv(data: bytes) -> tuple[Grid, list[str]]:
    warnings = []
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
        warnings.append("not UTF-8; read as Windows-1252 (cp1252)")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    try:
        rows = list(csv.reader(io.StringIO(text, newline=""), dialect))
    except csv.Error as exc:
        raise FileFormatError(f"not a readable CSV file ({exc})") from exc
    return Grid(None, rows), warnings


def read_xlsx(data: bytes, sheets: list[str] | None = None) -> list[Grid]:
    try:
        import openpyxl
    except ImportError as exc:
        raise FileFormatError("XLSX import needs openpyxl: uv pip install 'bmk[files]'") from exc
    try:
        # Not read-only: hyperlinks are only loaded in the full mode.
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # zipfile.BadZipFile, KeyError, InvalidFileException, …
        raise FileFormatError(f"not a readable XLSX file ({exc.__class__.__name__}: {exc})") from exc
    wanted = sheets or wb.sheetnames
    missing = [s for s in wanted if s not in wb.sheetnames]
    if missing:
        raise FileFormatError(f"no sheet {', '.join(missing)} (sheets: {', '.join(wb.sheetnames)})")
    grids = []
    for name in wanted:
        ws = wb[name]
        rows, links = [], {}
        for r, row in enumerate(ws.iter_rows()):
            rows.append([cell.value for cell in row])
            for c, cell in enumerate(row):
                if cell.hyperlink is not None and cell.hyperlink.target:
                    links[(r, c)] = cell.hyperlink.target
        while rows and not any(_text(v) for v in rows[-1]):
            rows.pop()
        grids.append(Grid(name, rows, links))
    return grids


def parse_table(
    data: bytes, fmt: str, *, url_col=None, title_col=None, date_col=None, sheets: list[str] | None = None, tz: str = "UTC"
) -> tuple[list[RawItem], list[str]]:
    if fmt == "xlsx":
        grids, warnings = read_xlsx(data, sheets), []
    else:
        grid, warnings = read_csv(data)
        grids = [grid]
    items = []
    for grid in grids:
        if not grid.rows:
            continue
        cols = detect(grid, url_col, title_col, date_col)
        got, notes = rows_to_items(grid, cols, tz)
        items += got
        warnings += notes
    return items, warnings


def preview(data: bytes, fmt: str, sheets: list[str] | None = None) -> list[dict[str, Any]]:
    """Detected columns and the first rows of each sheet, for the UI's column picker."""
    grids = read_xlsx(data, sheets) if fmt == "xlsx" else [read_csv(data)[0]]
    out = []
    for grid in grids:
        if not grid.rows:
            continue
        try:
            cols = detect(grid)
            out.append(
                {
                    "sheet": grid.name,
                    "headers": _headers(grid, cols.header),
                    "preview": [[_text(v) for v in row] for row in grid.rows[(1 if cols.header else 0) :][:PREVIEW_ROWS]],
                    "detected": {"url": cols.url, "title": cols.title, "date": cols.date, "header": cols.header},
                }
            )
        except ColumnsNeeded as exc:
            out.append({**exc.as_dict(), "detected": None})
    return out
